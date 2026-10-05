import csv
from io import BytesIO, StringIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class DatasetFileValidationError(ValueError):
    pass


def validate_dataset_file(filename: str, content: bytes) -> tuple[str, str, int, int]:
    extension = filename.rsplit(".", maxsplit=1)[-1].lower() if "." in filename else ""

    if not content:
        raise DatasetFileValidationError("The uploaded file is empty.")

    if extension == "csv":
        return "CSV", "text/csv", *_validate_csv(content)

    if extension == "xlsx":
        return (
            "Excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            *_validate_excel(content),
        )

    raise DatasetFileValidationError("Only CSV and .xlsx files are supported.")


def _validate_csv(content: bytes) -> tuple[int, int]:
    try:
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        header = next(reader, None)
        if not header or not any(cell.strip() for cell in header):
            raise DatasetFileValidationError("The CSV file must contain a header row.")

        column_count = len(header)
        row_count = 0
        for row in reader:
            if not row:
                continue
            if len(row) != column_count:
                raise DatasetFileValidationError("Every CSV row must have the same number of columns.")
            row_count += 1
        return row_count, column_count
    except UnicodeDecodeError as exc:
        raise DatasetFileValidationError("CSV files must use UTF-8 encoding.") from exc
    except csv.Error as exc:
        raise DatasetFileValidationError("The CSV file is malformed.") from exc


def _validate_excel(content: bytes) -> tuple[int, int]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise DatasetFileValidationError("The Excel file is invalid or corrupted.") from exc

    try:
        if not workbook.worksheets:
            raise DatasetFileValidationError("The Excel file must contain a worksheet.")

        row_count = 0
        column_count = 0
        has_header = False
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            populated_columns = [
                index
                for index, value in enumerate(row, start=1)
                if value is not None and (not isinstance(value, str) or value.strip())
            ]
            if not populated_columns:
                continue
            if not has_header:
                has_header = True
                column_count = max(populated_columns)
            else:
                row_count += 1

        if not has_header:
            raise DatasetFileValidationError("The Excel file must contain a header row.")
        return row_count, column_count
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, IndexError) as exc:
        raise DatasetFileValidationError("The Excel file is invalid or corrupted.") from exc
    finally:
        workbook.close()
