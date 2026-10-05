import csv
import json
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import Column, DateTime, Float, MetaData, String, Table, inspect
from sqlalchemy.orm import Session

from app.models.dataset import Dataset
from app.schemas.pipeline import PipelineRunCreate

MAX_RECORDED_ROW_ERRORS = 100


class PipelineExecutionError(ValueError):
    pass


@dataclass
class PipelineMetrics:
    total_rows: int
    valid_rows: int
    rejected_rows: int
    duplicates: int
    missing_values: int
    quality_score: float
    errors: list[str]


@dataclass
class PreparedDataset:
    headers: list[str]
    rows: list[dict[str, Any]]
    column_types: dict[str, Any]
    metrics: PipelineMetrics


def prepare_pipeline(
    dataset: Dataset,
    content: bytes,
    request: PipelineRunCreate,
) -> PreparedDataset:
    headers, source_rows = extract_dataset(dataset.file_type, content)
    return validate_clean_transform(headers, source_rows, request)


def extract_dataset(file_type: str, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    if file_type == "CSV":
        return _extract_csv(content)
    if file_type == "Excel":
        return _extract_excel(content)
    raise PipelineExecutionError(f"Unsupported dataset format: {file_type}.")


def _extract_csv(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        raw_headers = next(reader, None)
        if not raw_headers:
            raise PipelineExecutionError("The CSV file must contain a header row.")
        headers = _validated_headers(raw_headers)
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(reader, start=2):
            if not values:
                continue
            if len(values) != len(headers):
                raise PipelineExecutionError(
                    f"CSV row {row_number} has {len(values)} values; expected {len(headers)}."
                )
            rows.append(dict(zip(headers, values, strict=True)))
        return headers, rows
    except UnicodeDecodeError as exc:
        raise PipelineExecutionError("CSV files must use UTF-8 encoding.") from exc
    except csv.Error as exc:
        raise PipelineExecutionError("The CSV file is malformed.") from exc


def _extract_excel(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise PipelineExecutionError("The Excel file is invalid or corrupted.") from exc

    try:
        if not workbook.worksheets:
            raise PipelineExecutionError("The Excel file must contain a worksheet.")
        sheet = workbook.worksheets[0]
        iterator = iter(sheet.iter_rows(values_only=True))
        raw_headers: tuple[Any, ...] | None = None
        for row in iterator:
            if any(_is_populated(value) for value in row):
                raw_headers = row
                break
        if raw_headers is None:
            raise PipelineExecutionError("The Excel file must contain a header row.")
        last_header = max(index for index, value in enumerate(raw_headers) if _is_populated(value))
        headers = _validated_headers([str(value) if value is not None else "" for value in raw_headers[: last_header + 1]])
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(_is_populated(value) for value in values):
                continue
            if any(_is_populated(value) for value in values[len(headers) :]):
                raise PipelineExecutionError(
                    f"Excel row {row_number} has values beyond the {len(headers)} header columns."
                )
            padded_values = list(values[: len(headers)])
            padded_values.extend([None] * (len(headers) - len(padded_values)))
            rows.append(dict(zip(headers, padded_values, strict=True)))
        return headers, rows
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, IndexError) as exc:
        if isinstance(exc, PipelineExecutionError):
            raise
        raise PipelineExecutionError("The Excel file is invalid or corrupted.") from exc
    finally:
        workbook.close()


def _validated_headers(values: list[Any]) -> list[str]:
    headers = [str(value).strip() for value in values]
    if not headers or any(not header for header in headers):
        raise PipelineExecutionError("Every source column must have a non-empty header.")
    if len(headers) != len(set(headers)):
        raise PipelineExecutionError("Source column headers must be unique.")
    return headers


def validate_clean_transform(
    headers: list[str],
    source_rows: list[dict[str, Any]],
    request: PipelineRunCreate,
) -> PreparedDataset:
    selected_columns = set(request.numeric_columns) | set(request.date_columns)
    selected_columns |= set(request.drop_columns) | set(request.rename_columns)
    selected_columns |= set(request.lowercase_columns) | set(request.uppercase_columns)
    unknown_columns = selected_columns - set(headers)
    if unknown_columns:
        raise PipelineExecutionError(
            f"Unknown source column(s): {', '.join(sorted(unknown_columns))}."
        )
    if set(request.numeric_columns) & set(request.date_columns):
        raise PipelineExecutionError("A column cannot be both numeric and a date.")
    if set(request.lowercase_columns) & set(request.uppercase_columns):
        raise PipelineExecutionError("A column cannot be both lowercased and uppercased.")

    total_rows = len(source_rows)
    missing_values = sum(
        1
        for row in source_rows
        for value in row.values()
        if _is_missing(value)
    )
    cleaned_rows: list[dict[str, Any]] = []
    seen_rows: set[str] = set()
    duplicate_count = 0
    rejected_rows = 0
    row_errors: list[str] = []

    fill_values: dict[str, Any] = {}
    if request.missing_value_action == "fill":
        for column in request.numeric_columns:
            try:
                fill_values[column] = _parse_number(request.fill_value)
            except ValueError as exc:
                raise PipelineExecutionError(
                    f"Fill value is not numeric for column '{column}'."
                ) from exc
        for column in request.date_columns:
            try:
                fill_values[column] = _parse_date(request.fill_value)
            except ValueError as exc:
                raise PipelineExecutionError(
                    f"Fill value is not a valid date for column '{column}'."
                ) from exc

    for row_number, source_row in enumerate(source_rows, start=2):
        signature = json.dumps(source_row, sort_keys=True, default=str, ensure_ascii=True)
        if signature in seen_rows:
            duplicate_count += 1
            if request.duplicate_action == "drop":
                continue
        seen_rows.add(signature)

        row = dict(source_row)
        missing_columns = [column for column, value in row.items() if _is_missing(value)]
        if missing_columns and request.missing_value_action == "drop_row":
            rejected_rows += 1
            _record_error(row_errors, f"Row {row_number} rejected: missing {', '.join(missing_columns)}.")
            continue

        if request.missing_value_action == "fill":
            for column in missing_columns:
                row[column] = fill_values.get(column, request.fill_value)
        else:
            for column in missing_columns:
                row[column] = None

        invalid_values: list[str] = []
        for column in request.numeric_columns:
            if row[column] is not None:
                try:
                    row[column] = _parse_number(row[column])
                except ValueError:
                    invalid_values.append(f"{column} is not numeric")
        for column in request.date_columns:
            if row[column] is not None:
                try:
                    row[column] = _parse_date(row[column])
                except ValueError:
                    invalid_values.append(f"{column} is not a valid date")
        if invalid_values:
            rejected_rows += 1
            _record_error(row_errors, f"Row {row_number} rejected: {', '.join(invalid_values)}.")
            continue
        cleaned_rows.append(row)

    output_headers, output_rows, column_types = _transform_rows(headers, cleaned_rows, request)
    valid_rows = len(output_rows)
    quality_score = round((valid_rows / total_rows) * 100, 2) if total_rows else 0.0
    metrics = PipelineMetrics(
        total_rows=total_rows,
        valid_rows=valid_rows,
        rejected_rows=rejected_rows,
        duplicates=duplicate_count,
        missing_values=missing_values,
        quality_score=quality_score,
        errors=row_errors,
    )
    return PreparedDataset(output_headers, output_rows, column_types, metrics)


def _transform_rows(
    headers: list[str],
    rows: list[dict[str, Any]],
    request: PipelineRunCreate,
) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    kept_headers = [header for header in headers if header not in set(request.drop_columns)]
    renamed_headers = [
        request.rename_columns.get(header, header) for header in kept_headers
    ]
    if request.normalize_headers:
        output_headers = [_normalize_header(header, index) for index, header in enumerate(renamed_headers, 1)]
    else:
        output_headers = renamed_headers
    if len(output_headers) != len(set(output_headers)):
        raise PipelineExecutionError("Column transformations produced duplicate column names.")
    if any(len(header) > 63 for header in output_headers):
        raise PipelineExecutionError("Output column names must be 63 characters or fewer.")
    if not output_headers:
        raise PipelineExecutionError("At least one output column must remain after transformations.")

    output_rows: list[dict[str, Any]] = []
    for row in rows:
        output_row: dict[str, Any] = {}
        for source_header, output_header in zip(kept_headers, output_headers, strict=True):
            value = row[source_header]
            if isinstance(value, str):
                if request.trim_strings:
                    value = value.strip()
                if source_header in request.lowercase_columns:
                    value = value.lower()
                if source_header in request.uppercase_columns:
                    value = value.upper()
            output_row[output_header] = value
        output_rows.append(output_row)

    column_types: dict[str, Any] = {}
    for source_header, output_header in zip(kept_headers, output_headers, strict=True):
        if source_header in request.numeric_columns:
            column_types[output_header] = Float()
        elif source_header in request.date_columns:
            column_types[output_header] = DateTime(timezone=True)
        else:
            column_types[output_header] = String()
    return output_headers, output_rows, column_types


def load_rows(db: Session, destination_table: str, prepared: PreparedDataset) -> None:
    connection = db.connection()
    if inspect(connection).has_table(destination_table):
        raise PipelineExecutionError(
            f"Destination table '{destination_table}' already exists; choose a new table name."
        )
    metadata = MetaData()
    table = Table(
        destination_table,
        metadata,
        *[
            Column(column_name, prepared.column_types[column_name], nullable=True)
            for column_name in prepared.headers
        ],
    )
    table.create(connection)
    if prepared.rows:
        db.execute(table.insert(), prepared.rows)


def _is_populated(value: Any) -> bool:
    return not _is_missing(value)


def _is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_number(value: Any) -> float:
    if isinstance(value, bool):
        raise ValueError("Boolean values are not numeric.")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("Invalid numeric value.") from exc
    if not number.is_finite():
        raise ValueError("Numeric values must be finite.")
    result = float(number)
    if not math.isfinite(result):
        raise ValueError("Numeric value is out of range.")
    return result


def _parse_date(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    text = str(value).strip()
    if text.endswith("Z"):
        text = f"{text[:-1]}+00:00"
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for date_format in ("%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            continue
    raise ValueError("Invalid date value.")


def _normalize_header(value: str, index: int) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_]+", "_", value.strip()).strip("_").lower()
    return normalized or f"column_{index}"


def _record_error(errors: list[str], message: str) -> None:
    if len(errors) < MAX_RECORDED_ROW_ERRORS:
        errors.append(message)
    elif len(errors) == MAX_RECORDED_ROW_ERRORS:
        errors.append("Additional row errors were omitted.")
